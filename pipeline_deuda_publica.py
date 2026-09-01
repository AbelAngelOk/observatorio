"""
Pipeline de la deuda publica bruta (Secretaria de Finanzas).

Este indicador NO esta en la API de series: se publica solo como un boletin
mensual en Excel. Por eso es el pipeline mas fragil del proyecto -- depende del
layout de un .xlsx que puede cambiar sin aviso.

Confirmado contra el boletin real (boletin_mensual_31_05_2026.xlsx):
    * hoja "A.1" = Deuda Bruta de la Administracion Central, serie mensual 2019+;
    * la fila de fechas es la unica con muchas celdas de tipo fecha;
    * la fila del total arranca con "A- DEUDA BRUTA";
    * valores en millones de USD.
El parser autodetecta esas dos filas en vez de asumir numeros de fila fijos, que
es lo que hacia fallar a la version anterior.

La serie del boletin arranca en 2019 (es lo mas atras que publica la Secretaria
en este formato). Los puntos de 2023 que ya teniamos cargados a mano conviven:
el pipeline los revisa si el boletin trae un valor distinto para esa fecha.

Uso:
    pip install -r requirements.txt
    python pipeline_deuda_publica.py
"""

import datetime as dt
import re

import requests
import openpyxl

import core

NOMBRE = "deuda"  # como queda registrado en la tabla 'ingestas'
INDEX_PAGE = "https://www.argentina.gob.ar/economia/finanzas/datos-mensuales"
TMP_XLSX = "_tmp_deuda.xlsx"
UA = {"User-Agent": "Mozilla/5.0"}


def find_latest_excel_url() -> str:
    """Resuelve el link al ultimo boletin contra la pagina de indice. El nombre
    del archivo cambia todos los meses (boletin_mensual_31_MM_AAAA.xlsx), asi que
    no hay una URL estable: hay que leerla de la pagina."""
    resp = requests.get(INDEX_PAGE, timeout=30, headers=UA)
    resp.raise_for_status()
    # nos quedamos con los boletines mensuales, no con cualquier .xlsx suelto
    matches = re.findall(r'(https?://[^"\']+boletin_mensual[^"\']+\.xlsx)', resp.text, re.I)
    if not matches:
        matches = re.findall(r'href="([^"]+\.xlsx?)"', resp.text)
    if not matches:
        raise RuntimeError("No encontre ningun .xlsx en la pagina de indice; cambio el layout del sitio.")
    url = matches[0]
    if url.startswith("/"):
        url = "https://www.argentina.gob.ar" + url
    return url


def download_excel(url: str, dest: str) -> str:
    resp = requests.get(url, timeout=90, headers=UA)
    resp.raise_for_status()
    with open(dest, "wb") as f:
        f.write(resp.content)
    return dest


def parse_deuda_bruta(xlsx_path: str) -> list[dict]:
    """Serie mensual de deuda bruta total desde la hoja A.1.

    Autodetecta la fila de fechas (la que tiene mas celdas de tipo fecha) y la
    fila del total (la que arranca con "A- DEUDA BRUTA"), y las alinea por
    columna. Devuelve [{fecha, valor}] en millones de USD.
    """
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["A.1"] if "A.1" in wb.sheetnames else wb[wb.sheetnames[1]]

    fila_fechas = max(
        ws.iter_rows(min_row=1, max_row=20),
        key=lambda row: sum(1 for c in row if isinstance(c.value, dt.datetime)),
    )
    fechas = {c.column: c.value for c in fila_fechas if isinstance(c.value, dt.datetime)}
    if not fechas:
        raise RuntimeError("No encontre la fila de fechas en la hoja A.1; cambio el layout del boletin.")

    fila_total = None
    for row in ws.iter_rows(min_row=1, max_row=40):
        for c in row:
            if isinstance(c.value, str) and c.value.strip().upper().startswith("A- DEUDA BRUTA"):
                fila_total = row
                break
        if fila_total:
            break
    if fila_total is None:
        raise RuntimeError("No encontre la fila 'A- DEUDA BRUTA' en la hoja A.1.")

    series = []
    for c in fila_total:
        if c.column in fechas and isinstance(c.value, (int, float)):
            series.append({"fecha": fechas[c.column].strftime("%Y-%m-%d"),
                           "valor": round(float(c.value), 1)})
    series.sort(key=lambda r: r["fecha"])
    return series


def main() -> int:
    import os
    url = find_latest_excel_url()
    download_excel(url, TMP_XLSX)

    try:
        serie = parse_deuda_bruta(TMP_XLSX)
        with core.corrida(NOMBRE) as ctx:
            for punto in serie:
                ctx.nuevas += core.insertar_observacion(
                    ctx.conn, "deuda_publica_bruta", punto["fecha"], punto["valor"],
                    fuente_url=url, es_provisorio=False,
                )
            nuevas = ctx.nuevas
    finally:
        if os.path.exists(TMP_XLSX):
            os.remove(TMP_XLSX)

    print(f"OK deuda: {len(serie)} puntos en el boletin, {nuevas} nuevos o revisados.")
    return nuevas


if __name__ == "__main__":
    main()
