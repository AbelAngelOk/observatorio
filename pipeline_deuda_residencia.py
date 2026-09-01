"""
Pipeline experimental: deuda bruta de la Administracion Central por RESIDENCIA
del tenedor -- cuanto le debe el Estado a acreedores del exterior y cuanto a
acreedores locales.

Fuente: hoja A.4.5 ("Por residencia del tenedor") del boletin TRIMESTRAL de la
Secretaria de Finanzas. A diferencia de pipeline_cartera.py, que baja un archivo
por anio, aca alcanza con UN solo archivo: esa hoja trae la serie historica
completa (1994 -> hoy, ~117 puntos) en una sola tabla.

    deuda_no_residentes = columna "Deuda Externa"  (tenedores del exterior)
    deuda_residentes    = columna "Deuda Interna"  (tenedores locales)

Las dos suman el total y se validan contra la columna "Total Deuda".

QUE MIDE Y QUE NO -- importante antes de leer estas series:
  * Es deuda BRUTA de la Administracion Central. NO es la consolidada: no
    incluye los pasivos remunerados del BCRA (esos estan en deuda_consolidada).
  * NO es "deuda con terceros": la parte "residentes" incluye las tenencias
    de organismos del propio sector publico (ANSES/FGS, BCRA). Finanzas no
    publica el intra-sector publico como serie limpia, asi que no se puede
    restar -- es la misma limitacion documentada en deuda_consolidada.
  * La propia hoja EXCLUYE la deuda elegible pendiente de reestructuracion
    (holdouts), asi que su "Total Deuda" es algo menor que deuda_publica_bruta.
  * El corte por residencia es una estimacion de Finanzas construida sobre las
    cuentas internacionales del INDEC, no un censo de tenedores.

Unidades: la hoja publica en MILES DE MILLONES de USD; la base guarda millones
de USD (convencion del resto del proyecto), asi que se multiplica por 1000.

EXPERIMENTAL: depende del layout de un .xlsx que puede cambiar sin aviso. No
corre en el cron.

Uso:
    python pipeline_deuda_residencia.py
"""

import datetime as dt
import os
import re

import openpyxl
import requests

import core

NOMBRE = "deuda_residencia"
INDEX_PAGE = "https://www.argentina.gob.ar/economia/finanzas/datos-trimestrales-de-la-deuda"
TMP_XLSX = "_tmp_deuda_residencia.xlsx"
UA = {"User-Agent": "Mozilla/5.0"}

# La hoja publica en miles de millones de USD; la base usa millones.
A_MILLONES = 1000


def find_latest_excel_url() -> str:
    """URL del boletin trimestral mas reciente.

    Los archivos se llaman deuda_publica_DD-MM-AAAA.xlsx, asi que se elige por
    la FECHA del nombre y no por el orden en que aparecen en la pagina (que no
    esta garantizado)."""
    resp = requests.get(INDEX_PAGE, timeout=60, headers=UA)
    resp.raise_for_status()
    encontrados = {}
    for href in re.findall(r'href="([^"]+deuda_publica_[^"]+\.xlsx)"', resp.text, re.I):
        m = re.search(r"deuda_publica_(\d{2})-(\d{2})-(\d{4})", href)
        if not m:
            continue
        d, mth, y = m.groups()
        url = href if href.startswith("http") else "https://www.argentina.gob.ar" + href
        encontrados[f"{y}-{mth}-{d}"] = url
    if not encontrados:
        raise RuntimeError(
            "No encontre ningun deuda_publica_*.xlsx en la pagina de indice; cambio el layout.")
    ultima = max(encontrados)
    return encontrados[ultima]


def download_excel(url: str, dest: str) -> str:
    resp = requests.get(url, timeout=180, headers=UA)
    resp.raise_for_status()
    with open(dest, "wb") as f:
        f.write(resp.content)
    return dest


def _hoja_residencia(wb):
    """La hoja A.4.5. Si el codigo de hoja cambia, se la busca por su titulo
    ('POR RESIDENCIA DEL TENEDOR'), que es mas estable que el numero."""
    if "A.4.5" in wb.sheetnames:
        return wb["A.4.5"]
    for nombre in wb.sheetnames:
        ws = wb[nombre]
        for row in ws.iter_rows(min_row=1, max_row=15, max_col=6):
            for c in row:
                if isinstance(c.value, str) and "RESIDENCIA DEL TENEDOR" in c.value.upper():
                    return ws
    raise RuntimeError("No encontre la hoja de residencia del tenedor (A.4.5) en el boletin.")


def _columnas(ws):
    """Ubica las columnas de fecha / total / externa / interna leyendo la fila
    de encabezados, en vez de asumir que son B/C/D/E."""
    for row in ws.iter_rows(min_row=1, max_row=25, max_col=12):
        textos = {}
        for c in row:
            if isinstance(c.value, str):
                textos[c.value.strip().lower()] = c.column
        col_per = next((v for k, v in textos.items() if k.startswith("per")), None)
        col_tot = next((v for k, v in textos.items() if "total" in k), None)
        col_ext = next((v for k, v in textos.items() if "externa" in k and "%" not in k), None)
        col_int = next((v for k, v in textos.items() if "interna" in k and "%" not in k), None)
        if col_per and col_tot and col_ext and col_int:
            return row[0].row, col_per, col_tot, col_ext, col_int
    raise RuntimeError("No encontre la fila de encabezados (Periodo/Total/Externa/Interna) en A.4.5.")


def parse_residencia(xlsx_path: str) -> list[dict]:
    """[{fecha, externa, interna}] en millones de USD, ordenado por fecha.

    Saltea las filas sin apertura (la fuente trae 'n/d' en 2002-03) y valida
    que externa + interna reconstruya el total publicado."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = _hoja_residencia(wb)
    fila_head, col_per, col_tot, col_ext, col_int = _columnas(ws)

    puntos = []
    for row in ws.iter_rows(min_row=fila_head + 1, max_row=ws.max_row):
        celdas = {c.column: c.value for c in row}
        fecha = celdas.get(col_per)
        if not isinstance(fecha, dt.datetime):
            continue
        ext, inte, tot = celdas.get(col_ext), celdas.get(col_int), celdas.get(col_tot)
        # 'n/d' en las tres columnas: la fuente no publica la apertura ese trimestre.
        if not isinstance(ext, (int, float)) or not isinstance(inte, (int, float)):
            continue
        if isinstance(tot, (int, float)) and tot:
            if abs((ext + inte) - tot) / tot > 0.01:      # tolerancia 1%
                raise RuntimeError(
                    f"A.4.5 {fecha:%Y-%m-%d}: externa+interna ({ext + inte:.1f}) no reconstruye "
                    f"el total publicado ({tot:.1f}); cambio el significado de las columnas.")
        puntos.append({
            "fecha": fecha.strftime("%Y-%m-%d"),
            "externa": round(float(ext) * A_MILLONES, 1),
            "interna": round(float(inte) * A_MILLONES, 1),
        })
    if not puntos:
        raise RuntimeError("A.4.5 no devolvio ningun punto; cambio el layout de la hoja.")
    puntos.sort(key=lambda r: r["fecha"])
    return puntos


def main() -> int:
    url = find_latest_excel_url()
    download_excel(url, TMP_XLSX)
    try:
        serie = parse_residencia(TMP_XLSX)
        with core.corrida(NOMBRE) as ctx:
            for p in serie:
                ctx.nuevas += core.insertar_observacion(
                    ctx.conn, "deuda_no_residentes", p["fecha"], p["externa"], fuente_url=url)
                ctx.nuevas += core.insertar_observacion(
                    ctx.conn, "deuda_residentes", p["fecha"], p["interna"], fuente_url=url)
            nuevas = ctx.nuevas
    finally:
        if os.path.exists(TMP_XLSX):
            os.remove(TMP_XLSX)

    print(f"OK deuda_residencia: {len(serie)} trimestres ({serie[0]['fecha']} -> "
          f"{serie[-1]['fecha']}), {nuevas} observaciones nuevas o revisadas.")
    return nuevas


if __name__ == "__main__":
    main()
